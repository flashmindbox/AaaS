"""Build the AaaS judge pitch deck (pitch/AaaS-Pitch.pptx) - refined edition.

16:9, Odisha/Indian theme. Adds the open-source / Digital Public Infrastructure
angle, a flagship Indian Sign Language avatar slide, and a concrete "day in the
life" utility scenario. Diagrams are native editable shapes; Odia text, the
chakra and the sign avatar are pre-rendered images so they display anywhere.
Speaker notes on every slide. Footer auto-numbers.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import pathlib

REPO = pathlib.Path(__file__).resolve().parents[1]
AST = REPO / "pitch" / "assets"

CREAM = RGBColor(0xFF, 0xF7, 0xE9); PAPER = RGBColor(0xFD, 0xF3, 0xDD)
MAROON = RGBColor(0x8D, 0x1B, 0x2F); MAROON_D = RGBColor(0x6E, 0x13, 0x22)
SAFFRON = RGBColor(0xEF, 0x7D, 0x16); SAFF_L = RGBColor(0xFD, 0xEC, 0xD4)
GOLD = RGBColor(0xC9, 0x95, 0x2A); GOLD_B = RGBColor(0xE6, 0xB9, 0x3F)
INDIGO = RGBColor(0x1B, 0x2A, 0x66); IND_L = RGBColor(0xE9, 0xEC, 0xFB)
GREEN = RGBColor(0x1A, 0x7A, 0x3C); GRN_L = RGBColor(0xE4, 0xF4, 0xE9)
INK = RGBColor(0x3A, 0x24, 0x10); MUTED = RGBColor(0x7A, 0x5A, 0x38)
WHITE = RGBColor(0xFF, 0xFF, 0xFF); CREAMT = RGBColor(0xFF, 0xF3, 0xDA)
BLUET = RGBColor(0xBC, 0xD0, 0xFF)

FONT = "Segoe UI"; TOTAL = 19
prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
SW, SH = 13.333, 7.5; BLANK = prs.slide_layouts[6]


def slide():
    s = prs.slides.add_slide(BLANK); rect(s, 0, 0, SW, SH, CREAM); return s


def rect(s, x, y, w, h, fill, line=None, line_w=1.0, rounded=False, radius=0.09, shadow=False):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                             Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        try: shp.adjustments[0] = radius
        except Exception: pass
    if fill is None: shp.fill.background()
    else: shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb = line; shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    if shadow:
        el = shp._element.spPr
        ef = el.makeelement(qn('a:effectLst'), {}); el.append(ef)
        sh = ef.makeelement(qn('a:outerShdw'), {'blurRad':'90000','dist':'40000','dir':'5400000','rotWithShape':'0'}); ef.append(sh)
        c = sh.makeelement(qn('a:srgbClr'), {'val':'6E1322'}); sh.append(c)
        a = c.makeelement(qn('a:alpha'), {'val':'22000'}); c.append(a)
    return shp


def txt(s, x, y, w, h, runs, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, font=FONT, spacing=1.0, italic=False):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    items = runs if isinstance(runs, list) else [runs]
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing: p.line_spacing = spacing
        p.space_after = Pt(4)
        text, opts = it if isinstance(it, tuple) else (it, {})
        r = p.add_run(); r.text = text; f = r.font
        f.size = Pt(opts.get("size", size)); f.bold = opts.get("bold", bold)
        f.italic = opts.get("italic", italic); f.name = opts.get("font", font)
        f.color.rgb = opts.get("color", color)
    return tb


def bullets(s, x, y, w, h, items, size=18, color=INK, gap=6, marker="▸  ", mcolor=SAFFRON):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.05; p.space_after = Pt(gap)
        head = it[0] if isinstance(it, tuple) else None
        body = it[1] if isinstance(it, tuple) else it
        rm = p.add_run(); rm.text = marker
        rm.font.size = Pt(size); rm.font.bold = True; rm.font.name = FONT; rm.font.color.rgb = mcolor
        if head:
            rh = p.add_run(); rh.text = head + "  "
            rh.font.size = Pt(size); rh.font.bold = True; rh.font.name = FONT; rh.font.color.rgb = INK
        rb = p.add_run(); rb.text = body
        rb.font.size = Pt(size); rb.font.name = FONT; rb.font.color.rgb = color
    return tb


def arrow(s, x1, y1, x2, y2, color=GOLD, w=2.0):
    cn = s.shapes.add_connector(2, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    cn.line.color.rgb = color; cn.line.width = Pt(w)
    ln = cn.line._get_or_add_ln()
    ln.append(ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'}))
    cn.shadow.inherit = False; return cn


def pic(s, path, x, y, w=None, h=None):
    kw = {}
    if w: kw["width"] = Inches(w)
    if h: kw["height"] = Inches(h)
    return s.shapes.add_picture(str(path), Inches(x), Inches(y), **kw)


def title_bar(s, title, kicker=None):
    rect(s, 0, 0, SW, 1.0, MAROON); rect(s, 0, 1.0, SW, 0.07, GOLD_B)
    txt(s, 0.55, 0.0, 10.8, 1.0, title, size=29, color=CREAMT, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    if kicker:
        txt(s, 0.57, 0.63, 11.5, 0.32, kicker, size=13.5, color=GOLD_B, bold=True)
    pic(s, AST / "chakra_gold.png", 12.35, 0.13, h=0.74)


def footer(s):
    n = len(prs.slides._sldIdLst)
    txt(s, 0.55, 7.06, 8.0, 0.35, "Team SUBARNAREKHA  ·  Smart Odisha Hackathon '25", size=10, color=MUTED)
    txt(s, 11.2, 7.06, 1.6, 0.35, f"{n:02d} / {TOTAL}", size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def notes(s, text):
    s.notes_slide.notes_text_frame.text = text


def box(s, x, y, w, h, title, sub, fill, border, tcolor=INK, scolor=MUTED, tsize=15, ssize=11):
    rect(s, x, y, w, h, fill, line=border, line_w=1.75, rounded=True, radius=0.12)
    tb = s.shapes.add_textbox(Inches(x+0.12), Inches(y), Inches(w-0.24), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.05); tf.margin_right = Inches(0.05)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; p.line_spacing = 1.0
    r = p.add_run(); r.text = title; r.font.size = Pt(tsize); r.font.bold = True; r.font.name = FONT; r.font.color.rgb = tcolor
    if sub:
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER; p2.line_spacing = 1.0; p2.space_before = Pt(2)
        r2 = p2.add_run(); r2.text = sub; r2.font.size = Pt(ssize); r2.font.name = FONT; r2.font.color.rgb = scolor


def pill(s, x, y, w, h, text, fill, tcolor, size=13):
    rect(s, x, y, w, h, fill, rounded=True, radius=0.5)
    txt(s, x, y, w, h, text, size=size, color=tcolor, bold=True, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)


# ============================================================ 1 TITLE
s = slide()
rect(s, 0, 0, SW, SH, MAROON); rect(s, 0, 0, SW, 0.14, GOLD_B); rect(s, 0, SH-0.14, SW, 0.14, GOLD_B)
pic(s, AST / "chakra_gold.png", 0.7, 1.35, h=3.3)
txt(s, 4.5, 1.4, 8.4, 1.2, "Accessibility as a Service", size=50, color=CREAMT, bold=True)
txt(s, 4.52, 2.42, 8.2, 0.6, "(AaaS)", size=25, color=GOLD_B, bold=True)
pic(s, AST / "odia_sub.png", 4.52, 3.02, h=0.58)
txt(s, 4.52, 3.72, 8.4, 0.6, "Shared accessibility infrastructure for public institutions", size=19, color=CREAM)
pill(s, 4.55, 4.42, 5.6, 0.5, "Open-source · Digital Public Infrastructure", GOLD_B, MAROON_D, size=14)
txt(s, 4.5, 5.15, 8.4, 1.4, [
    ("Team SUBARNAREKHA", {"size": 23, "bold": True, "color": WHITE}),
    ("Smart Odisha Hackathon '25  ·  Software · Odiapreneur track", {"size": 15, "color": GOLD_B}),
], spacing=1.15)
notes(s, "Namaskar. We are Team SUBARNAREKHA. AaaS is open-source accessibility infrastructure - "
          "shared rails that make any public website usable by every citizen, in their own "
          "language, including the blind, deaf, and dyslexic. One line for institutions; one "
          "extension for citizens.")

# ============================================================ 2 PROBLEM
s = slide(); title_bar(s, "The Problem")
bullets(s, 0.7, 1.4, 12.0, 5.2, [
    ("26.8M+ Indians live with a disability", "(Census 2011) - most government & education sites are unusable with a screen reader."),
    ("The web is English-first.", "Only ~1 in 10 Indians reads English comfortably - Odia-only citizens are shut out."),
    ("Deaf citizens are invisible online.", "Indian Sign Language is their first language, yet there is almost no public content in it."),
    ("Today's \"fix\" is an overlay sold to site owners.", "English-centric, works only where installed, criticised for not fixing real access."),
    ("Every department reinvents accessibility.", "Slow, costly, inconsistent - and rarely WCAG / GIGW compliant."),
], size=18, gap=12)
rect(s, 0.7, 6.5, 12.0, 0.02, GOLD)
txt(s, 0.7, 6.58, 12.0, 0.5,
    "If you cannot see the screen, cannot read English, or cannot hear - public services are closed to you.",
    size=14.5, color=MAROON, bold=True, italic=True)
footer(s)
notes(s, "Three exclusions - sight, language, hearing - plus no shared solution. The deaf/ISL line "
          "sets up our flagship later. Overlays are English band-aids.")

# ============================================================ 3 WHO
s = slide(); title_bar(s, "Who gets left behind")
cards = [
    ("The blind citizen", "A scholarship form they literally cannot read - no audio, no screen reader.", MAROON, SAFF_L),
    ("The Odia speaker", "Fluent in Odia, facing an English-only portal with no translation.", GREEN, GRN_L),
    ("The exam candidate", "Dyslexia, in an online test - no audio, no extra time, no friendly fonts.", INDIGO, IND_L),
    ("The deaf citizen", "Misses every notice - nothing is offered in Indian Sign Language.", MAROON_D, SAFF_L),
]
cw, gap, x0 = 2.8, 0.27, 0.7
for i, (t, d, accent, fill) in enumerate(cards):
    x = x0 + i * (cw + gap)
    rect(s, x, 1.55, cw, 4.5, fill, line=accent, line_w=2, rounded=True, radius=0.06, shadow=True)
    rect(s, x, 1.55, cw, 0.65, accent, rounded=True, radius=0.2); rect(s, x, 1.9, cw, 0.3, accent)
    txt(s, x+0.15, 1.55, cw-0.3, 0.65, t, size=15.5, color=CREAMT, bold=True, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    txt(s, x+0.22, 2.45, cw-0.44, 3.4, d, size=14, color=INK, spacing=1.12)
txt(s, 0.7, 6.4, 12.0, 0.6,
    "Not edge cases - millions of citizens public services are meant to serve.",
    size=15, color=MAROON, bold=True, italic=True)
footer(s)
notes(s, "Four faces, including the deaf citizen - foreshadows the sign-language avatar. Land on "
          "scale: millions.")

# ============================================================ 4 SOLUTION
s = slide(); title_bar(s, "Our Solution: one shared platform")
rect(s, 0.7, 1.4, 12.0, 1.05, MAROON, rounded=True, radius=0.12)
txt(s, 0.9, 1.4, 11.6, 1.05, "Institutions plug in ONCE; citizens get accessibility EVERYWHERE.",
    size=22, color=CREAMT, bold=True, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
pillars = [
    ("API + Widget", "Integrate once", "One <script> line or a REST API turns any site accessible.", SAFFRON, SAFF_L),
    ("Browser Extension", "Works on any site", "Citizens install once - accessibility on every website, no site cooperation.", GREEN, GRN_L),
    ("On-device Indian AI", "Private & offline", "Neural Odia/Hindi/English voice & translation run in the browser.", INDIGO, IND_L),
]
cw = 3.95
for i, (t, tag, d, accent, fill) in enumerate(pillars):
    x = 0.7 + i * (cw + 0.30)
    rect(s, x, 2.65, cw, 3.25, fill, line=accent, line_w=2, rounded=True, radius=0.07, shadow=True)
    txt(s, x+0.28, 2.85, cw-0.56, 0.6, t, size=19, color=accent, bold=True)
    txt(s, x+0.28, 3.45, cw-0.56, 0.4, tag.upper(), size=12, color=MUTED, bold=True)
    txt(s, x+0.28, 3.95, cw-0.56, 1.8, d, size=15, color=INK, spacing=1.12)
rect(s, 0.7, 6.15, 12.0, 0.78, GOLD_B, rounded=True, radius=0.3)
txt(s, 0.7, 6.15, 12.0, 0.78,
    "Open-source · built on open AI (Meta MMS · Whisper · AI4Bharat) · no vendor lock-in",
    size=16, color=MAROON_D, bold=True, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
footer(s)
notes(s, "One platform, not a point tool. Three pillars + the open-source ribbon: we stand on open "
          "Indian/global AI, and the platform itself is open and self-hostable.")

# ============================================================ 5 TWO WAYS IN
s = slide(); title_bar(s, "Two ways in - the same powers")
data = [
    ("Institutions integrate", INDIGO, IND_L, [
        "Add one widget <script> line, or call the REST API.",
        "Site gains read-aloud, translation, voice-to-fill forms and an accessible exam mode.",
        "Central control: API keys, usage metering, WCAG logging.",
    ]),
    ("Citizens install the extension", GREEN, GRN_L, [
        "One-time install in Chrome / Edge.",
        "Accessibility on ANY website - including legacy portals that will never be updated.",
        "Runs on-device: offline and even on strict-CSP sites.",
    ]),
]
for col, (t, accent, fill, items) in enumerate(data):
    x = 0.7 + col * 6.25
    rect(s, x, 1.55, 5.9, 4.4, fill, line=accent, line_w=2, rounded=True, radius=0.06, shadow=True)
    rect(s, x, 1.55, 5.9, 0.75, accent, rounded=True, radius=0.16); rect(s, x, 1.95, 5.9, 0.35, accent)
    txt(s, x+0.3, 1.55, 5.3, 0.75, t, size=19, color=CREAMT, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    bullets(s, x+0.32, 2.55, 5.3, 3.3, items, size=15, gap=10, mcolor=accent)
pill(s, 0.7, 6.35, 12.0, 0.7, "Power with the user - not just the site owner.", GOLD_B, MAROON_D, size=18)
footer(s)
notes(s, "This pre-empts the overlay question. Overlays are site-owner-only; our extension gives a "
          "citizen-controlled path that reaches sites that will never integrate.")

# ============================================================ 6 ARCHITECTURE
s = slide(); title_bar(s, "Core architecture", "two ways in -> one API gateway -> shared engines (also run on-device)")
for lx, lab in ((1.9, "CLIENTS"), (6.55, "API GATEWAY"), (10.95, "ENGINES")):
    txt(s, lx-1.4, 1.5, 2.8, 0.3, lab, size=12, color=MAROON, bold=True, align=PP_ALIGN.CENTER)
box(s, 0.7, 1.95, 2.95, 1.35, "Institution site / app", "widget (1 line) + REST API", SAFF_L, SAFFRON)
box(s, 0.7, 3.55, 2.95, 1.35, "Browser extension", "any website, on-device", SAFF_L, SAFFRON)
rect(s, 4.35, 1.95, 4.05, 2.95, MAROON, line=GOLD_B, line_w=2, rounded=True, radius=0.05, shadow=True)
txt(s, 4.45, 2.05, 3.85, 0.5, "Accessibility API Gateway", size=17, color=CREAMT, bold=True, align=PP_ALIGN.CENTER)
txt(s, 4.45, 2.5, 3.85, 0.3, "X-API-Key · multi-tenant · CORS", size=11.5, color=GOLD_B, align=PP_ALIGN.CENTER)
for i, ep in enumerate(["POST /tts", "POST /stt", "POST /translate", "GET /widget.js"]):
    cx = 4.5 + (i % 2) * 1.95; cy = 2.95 + (i // 2) * 0.5
    rect(s, cx, cy, 1.8, 0.4, CREAMT, rounded=True, radius=0.3)
    txt(s, cx, cy, 1.8, 0.4, ep, size=11, color=MAROON_D, bold=True, font="Consolas", align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
txt(s, 4.45, 4.05, 3.85, 0.3, "usage metering · rate limiting · WCAG logging", size=11, color=GOLD_B, align=PP_ALIGN.CENTER)
box(s, 4.35, 5.05, 4.05, 0.95, "Admin Dashboard", "tenants · API keys · usage · WCAG", INDIGO, GOLD_B, tcolor=CREAMT, scolor=BLUET)
arrow(s, 4.35, 5.05, 6.37, 4.95, GOLD)
box(s, 9.05, 1.95, 3.55, 0.85, "Text-to-Speech + Screen Reader", "Meta MMS-TTS · Odia/Hindi/English", IND_L, INDIGO, tsize=13.5, ssize=11)
box(s, 9.05, 2.98, 3.55, 0.85, "Speech-to-Text", "Whisper", SAFF_L, SAFFRON, tsize=13.5, ssize=11)
box(s, 9.05, 4.01, 3.55, 0.85, "Translation", "AI4Bharat IndicTrans2", GRN_L, GREEN, tsize=13.5, ssize=11)
arrow(s, 3.65, 2.62, 4.35, 2.9, GOLD); arrow(s, 3.65, 4.22, 4.35, 3.6, GOLD)
arrow(s, 8.4, 2.7, 9.05, 2.38, GOLD); arrow(s, 8.4, 3.2, 9.05, 3.4, GOLD); arrow(s, 8.4, 3.7, 9.05, 4.43, GOLD)
txt(s, 0.7, 6.25, 12.0, 0.7,
    "On-device option: the extension runs all three engines in the browser via WebAssembly - offline, even on strict-CSP sites.",
    size=13.5, color=MUTED, italic=True)
footer(s)
notes(s, "Left to right. The gateway is the product: secure, multi-tenant, metered. Same engines "
          "run server-side OR on-device. Admin dashboard gives departments WCAG evidence.")

# ============================================================ 7 HOW IT WORKS
s = slide(); title_bar(s, "How it works")
rows = [
    ("Read aloud", INDIGO, IND_L, ["Page text", "IndicTrans2 (if needed)", "MMS-TTS on-device", "Odia voice plays"]),
    ("Translate page", GREEN, GRN_L, ["Page text", "IndicTrans2", "Page rewritten in Odia"]),
    ("Speak to fill", SAFFRON, SAFF_L, ["Voice (mic)", "Whisper on-device", "Form field auto-filled"]),
]
y = 1.65
for name, accent, fill, steps in rows:
    rect(s, 0.7, y, 2.3, 0.92, accent, rounded=True, radius=0.12)
    txt(s, 0.75, y, 2.2, 0.92, name, size=16, color=CREAMT, bold=True, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    sx = 3.25; sw = (12.6 - sx) / len(steps) - 0.25
    for i, st in enumerate(steps):
        x = sx + i * (sw + 0.25)
        rect(s, x, y+0.11, sw, 0.7, fill, line=accent, line_w=1.5, rounded=True, radius=0.12)
        txt(s, x+0.05, y+0.11, sw-0.1, 0.7, st, size=12.5, color=INK, bold=True, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
        if i < len(steps)-1: arrow(s, x+sw, y+0.46, x+sw+0.25, y+0.46, accent, 1.75)
    y += 1.3
rect(s, 0.7, 6.15, 12.0, 0.82, PAPER, line=GOLD, line_w=1.25, rounded=True, radius=0.2)
txt(s, 0.95, 6.15, 11.5, 0.82,
    "\"On-device\" = runs inside the browser via WebAssembly. No server, fully offline, and translation is skipped when the page is already in your language.",
    size=13.5, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
footer(s)
notes(s, "Three everyday actions, each a tiny pipeline. The heavy AI runs on the user's device - "
          "no backend, privacy by design.")

# ============================================================ 8 A DAY WITH AaaS
s = slide(); title_bar(s, "A day with AaaS", "Sunita applies for a scholarship on an English-only portal")
steps = [
    ("Opens the portal and taps the AaaS button - no app store, no setup.", INDIGO),
    ("The page reads aloud in Odia; she hovers to hear each field.", GREEN),
    ("She fills the form by voice - her speech lands in the right fields.", SAFFRON),
    ("Her deaf brother watches the same notice in Indian Sign Language, via the avatar.", MAROON),
]
y0 = 1.7
for i, (t, accent) in enumerate(steps):
    cy = y0 + i * 1.0
    rect(s, 0.9, cy, 0.72, 0.72, accent, rounded=True, radius=0.5)
    txt(s, 0.9, cy, 0.72, 0.72, str(i+1), size=26, color=CREAMT, bold=True, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    if i < len(steps)-1: arrow(s, 1.26, cy+0.72, 1.26, cy+1.0, GOLD, 2.5)
    rect(s, 1.95, cy+0.04, 10.6, 0.64, PAPER, line=GOLD, line_w=1, rounded=True, radius=0.12)
    txt(s, 2.2, cy+0.04, 10.1, 0.64, t, size=16, color=INK, anchor=MSO_ANCHOR.MIDDLE)
pill(s, 0.9, 6.05, 11.65, 0.78,
     "No sighted help. No English. No department changing a single line of code.",
     MAROON, CREAMT, size=17)
footer(s)
notes(s, "Make utility concrete - one citizen, one journey, touching all four powers (read, "
          "translate, voice, sign). The closing line is the payoff: zero dependence on the site.")

# ============================================================ 9 DEMO
s = slide(); title_bar(s, "It already works - live, on any site")
pic(s, AST / "demo-wikipedia.png", 0.7, 1.45, w=8.7)
rect(s, 0.7, 1.45, 8.7, 4.13, None, line=GOLD, line_w=1.5)
txt(s, 9.65, 1.7, 3.2, 4.4, [
    ("English Wikipedia -> Odia", {"size": 18, "bold": True, "color": MAROON}),
    ("Title and page flip to Odia in place, and the page reads aloud in a neural Odia voice.", {"size": 14, "color": INK}),
    ("No backend running.", {"size": 14, "bold": True, "color": GREEN}),
    ("Works despite Wikipedia's strict Content-Security-Policy - translation goes through the extension, the voice runs on-device.", {"size": 14, "color": INK}),
], spacing=1.12)
footer(s)
notes(s, "A real screenshot, not a mockup. Wikipedia is a deliberately hard target (strict CSP). "
          "If it works here, it works on government portals. Offer to demo live.")

# ============================================================ 10 SIGN LANGUAGE (FLAGSHIP)
s = slide(); title_bar(s, "Flagship: Indian Sign Language avatar")
rect(s, 0.7, 1.5, 4.4, 5.0, MAROON, rounded=True, radius=0.05, shadow=True)
pic(s, AST / "signavatar.png", 1.55, 2.0, h=2.9)
pill(s, 1.35, 5.05, 3.1, 0.55, "In development - our north star", GOLD_B, MAROON_D, size=13)
txt(s, 0.9, 5.75, 4.0, 0.7, "Turning every notice into sign language - at scale.",
    size=15, color=CREAMT, bold=True, italic=True, align=PP_ALIGN.CENTER)
bullets(s, 5.4, 1.7, 7.3, 4.6, [
    ("First language, not a translation:", "for millions of deaf Indians, Indian Sign Language (ISL) is primary - yet public services offer almost nothing in it."),
    ("What it does:", "converts on-screen text and announcements into ISL, performed by a 3D avatar - reusing the government's ISLRTC ISL dictionary."),
    ("How:", "text -> ISL gloss -> avatar animation, delivered through the same widget & extension."),
    ("Why it's a leap:", "no overlay or mainstream tool offers sign language - this reaches a community everyone else ignores."),
], size=16, gap=14)
footer(s)
notes(s, "Our wow + heart slide. Be honest: it's in development - the flagship next step. ISL is "
          "the deaf community's first language; building on ISLRTC's official dictionary makes it "
          "credible and India-specific. No competitor does this.")

# ============================================================ 11 OPEN SOURCE / DPI
s = slide(); title_bar(s, "Open-source: accessibility as public infrastructure")
rect(s, 0.7, 1.4, 12.0, 1.0, INDIGO, rounded=True, radius=0.1)
txt(s, 0.9, 1.4, 11.6, 1.0, "Like UPI and DigiLocker - but for accessibility: open rails any institution can build on.",
    size=20, color=CREAMT, bold=True, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
cards = [
    ("Open by default", "Permissive license, auditable code - essential for government security & privacy review.", GREEN, GRN_L),
    ("Built on open AI", "Meta MMS-TTS, OpenAI Whisper, AI4Bharat IndicTrans2 - no black boxes, no foreign lock-in.", SAFFRON, SAFF_L),
    ("Community-powered", "Anyone can add languages, voices and signs - it grows with India.", INDIGO, IND_L),
    ("Public good", "Free for institutions to deploy, sovereign and self-hostable - no recurring vendor fees.", MAROON, SAFF_L),
]
cw = 2.95
for i, (t, d, accent, fill) in enumerate(cards):
    x = 0.7 + i * (cw + 0.13)
    rect(s, x, 2.65, cw, 3.55, fill, line=accent, line_w=2, rounded=True, radius=0.06, shadow=True)
    rect(s, x, 2.65, cw, 0.62, accent, rounded=True, radius=0.18); rect(s, x, 2.95, cw, 0.32, accent)
    txt(s, x+0.15, 2.65, cw-0.3, 0.62, t, size=15, color=CREAMT, bold=True, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    txt(s, x+0.22, 3.45, cw-0.44, 2.6, d, size=14, color=INK, spacing=1.12)
footer(s)
notes(s, "The DPI framing govt judges love - reusable public rails, like UPI. Open-source means "
          "auditable (security), sovereign (no foreign lock-in), and community-extensible "
          "(more languages and signs). It also lets the state own its accessibility stack.")

# ============================================================ 12 SIX PARTS
s = slide(); title_bar(s, "One platform, six parts")
parts = [
    ("1", "Speech-to-Text", "Voice input, dictation", SAFFRON),
    ("2", "Text-to-Speech + Screen Reader", "Neural Odia/Hindi/English", INDIGO),
    ("3", "Accessibility API Gateway", "The plug-and-play surface", MAROON),
    ("4", "Translation - IndicTrans2", "Indic-tuned, Odia-first", GREEN),
    ("5", "Accessible Exam module", "Extra time, audio, dyslexia fonts", SAFFRON),
    ("6", "Admin Dashboard", "Tenants, keys, usage, WCAG", INDIGO),
]
cw, ch = 3.95, 1.55
for i, (n, t, d, accent) in enumerate(parts):
    x = 0.7 + (i % 3) * (cw + 0.27); y = 1.7 + (i // 3) * (ch + 0.3)
    rect(s, x, y, cw, ch, PAPER, line=GOLD, line_w=1.25, rounded=True, radius=0.08, shadow=True)
    rect(s, x+0.2, y+0.3, 0.95, 0.95, accent, rounded=True, radius=0.5)
    txt(s, x+0.2, y+0.3, 0.95, 0.95, n, size=30, color=CREAMT, bold=True, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    txt(s, x+1.35, y+0.32, cw-1.55, 0.95, [
        (t, {"size": 15.5, "bold": True, "color": INK}), (d, {"size": 12.5, "color": MUTED})], spacing=1.05)
footer(s)
notes(s, "Maps directly to the problem statement's six modules - we built the whole platform, with "
          "the exam module and admin oversight as differentiators.")

# ============================================================ 13 TECH
s = slide(); title_bar(s, "Built to run anywhere")
bullets(s, 0.7, 1.5, 7.4, 5.2, [
    ("Browser extension (MV3):", "MAIN + isolated worlds, a Shadow-DOM widget that can't clash with any page."),
    ("FastAPI gateway:", "multi-tenant, X-API-Key auth, CORS, usage metering, WCAG logging."),
    ("Open AI models:", "Meta MMS-TTS (or/hi/en), OpenAI Whisper, AI4Bharat IndicTrans2."),
    ("On-device runtime:", "transformers.js + ONNX / WebAssembly - inference in the browser."),
    ("Open-source platform:", "permissive license, self-hostable; runs offline from a USB / folder."),
], size=16.5, gap=12)
rect(s, 8.4, 1.6, 4.25, 4.9, MAROON, rounded=True, radius=0.06, shadow=True)
txt(s, 8.7, 1.85, 3.7, 0.5, "Why it matters", size=18, color=GOLD_B, bold=True)
bullets(s, 8.7, 2.55, 3.7, 3.8, [
    "Privacy: voice & text stay on the device.",
    "Offline: no connectivity required.",
    "Resilient: works on strict-CSP sites.",
    "Auditable: open code for security review.",
    "Low cost: no per-call cloud AI bills.",
], size=14, gap=11, mcolor=GOLD_B, color=CREAM)
footer(s)
notes(s, "For technical judges. Headline win: on-device inference - privacy, offline, resilience, "
          "auditability and cost all at once. Everything sits on open models + an open platform.")

# ============================================================ 14 OVERLAY
s = slide(); title_bar(s, "Isn't this just another accessibility overlay?")
rows = [
    ("", "Overlays (accessiBe / UserWay)", "AaaS"),
    ("Who controls it", "Site owner installs & pays; one site only", "User installs the extension -> works on ANY site"),
    ("Language", "English / Western-centric", "Odia & Hindi first; sign language on the way"),
    ("Data", "Cloud SaaS; content sent to their servers", "On-device & in-country; DPDP-friendly"),
    ("Openness", "Proprietary black box", "Open-source, auditable, self-hostable"),
    ("Claim", "\"Auto-makes you compliant\" (criticised)", "Honest: assistive layer + API toolkit"),
]
tw, th = 12.0, 4.7
tbl = s.shapes.add_table(len(rows), 3, Inches(0.7), Inches(1.5), Inches(tw), Inches(th)).table
tbl.columns[0].width = Inches(2.4); tbl.columns[1].width = Inches(4.8); tbl.columns[2].width = Inches(4.8)
for r, row in enumerate(rows):
    tbl.rows[r].height = Inches(th/len(rows))
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.margin_left = Inches(0.12); cell.margin_right = Inches(0.1); cell.margin_top = Inches(0.03); cell.margin_bottom = Inches(0.03)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]; cell.text_frame.word_wrap = True
        rn = p.add_run(); rn.text = val; rn.font.name = FONT
        if r == 0:
            cell.fill.solid(); cell.fill.fore_color.rgb = MAROON if c != 2 else GREEN
            rn.font.color.rgb = CREAMT; rn.font.bold = True; rn.font.size = Pt(14)
        else:
            cell.fill.solid(); cell.fill.fore_color.rgb = PAPER if c != 2 else GRN_L
            rn.font.size = Pt(12.5); rn.font.color.rgb = INK
            if c == 0: rn.font.bold = True; rn.font.color.rgb = MAROON
            if c == 2: rn.font.bold = True; rn.font.color.rgb = RGBColor(0x0F, 0x5A, 0x2C)
txt(s, 0.7, 6.4, 12.0, 0.5,
    "Not a compliance shortcut - open accessibility infrastructure with first-class Indian-language access.",
    size=14, color=MAROON, bold=True, italic=True)
footer(s)
notes(s, "Answer the hard question head-on. Overlays: site-owner, English, cloud, proprietary, "
          "over-claim compliance. We: user-controlled, Indian-language, on-device, open-source, honest.")

# ============================================================ 15 WHY GOVT
s = slide(); title_bar(s, "Why government should adopt AaaS")
bullets(s, 0.7, 1.45, 12.0, 5.4, [
    ("Compliance, built in:", "helps meet the RPwD Act 2016, GIGW 3.0 and WCAG 2.1 - with WCAG logging for evidence."),
    ("Real inclusion:", "first-class Odia & Hindi, and Indian Sign Language - reaching citizens English-only tools miss."),
    ("Data sovereignty:", "on-device & in-country processing - DPDP Act 2023 friendly; no citizen data to foreign clouds."),
    ("Open-source & no lock-in:", "auditable, self-hostable, sovereign - the state owns its accessibility stack, no vendor fees."),
    ("Cost-efficient & instant:", "one shared platform vs every dept reinventing it - and it works on existing legacy portals."),
], size=17.5, gap=13)
footer(s)
notes(s, "Procurement language: compliance, sovereignty, open-source/no-lock-in, cost, and reach. "
          "Open-source is the clincher - government can audit and own it.")

# ============================================================ 16 IMPACT
s = slide(); title_bar(s, "Impact & alignment")
rect(s, 0.7, 1.55, 5.9, 4.9, PAPER, line=GOLD, line_w=1.25, rounded=True, radius=0.05, shadow=True)
txt(s, 0.95, 1.75, 5.4, 0.5, "Outcomes", size=18, color=MAROON, bold=True)
bullets(s, 0.95, 2.4, 5.4, 3.8, [
    "Any citizen can hear & read any public page in Odia.",
    "Sign-language access for the deaf community.",
    "Exams become accessible for students with disabilities.",
    "Measurable WCAG compliance - no site rebuild.",
], size=15.5, gap=12)
rect(s, 6.75, 1.55, 5.9, 4.9, MAROON, rounded=True, radius=0.05, shadow=True)
txt(s, 7.0, 1.75, 5.4, 0.5, "Aligns with national missions", size=18, color=GOLD_B, bold=True)
bullets(s, 7.0, 2.4, 5.4, 3.8, [
    "Digital India - inclusive e-governance.",
    "Sugamya Bharat (Accessible India Campaign).",
    "Digital Public Infrastructure - open rails.",
    "SDG 10 & SDG 4 - equity and education.",
], size=15.5, gap=12, mcolor=GOLD_B, color=CREAM)
footer(s)
notes(s, "Tie outcomes to flagship missions - now including DPI. Signals policy alignment, not a toy.")

# ============================================================ 17 ROADMAP
s = slide(); title_bar(s, "Way forward")
phases = [
    ("NOW - MVP", GREEN, GRN_L, ["Extension + gateway live", "Odia / Hindi / English", "On-device TTS & STT", "Open-source + offline bundle"]),
    ("NEXT - 3-6 months", SAFFRON, SAFF_L, ["Indian Sign Language avatar *", "More Indian languages", "Pilot with a govt department", "Android / mobile"]),
    ("LATER - 6-12 months", INDIGO, IND_L, ["State-wide DPI rollout", "WCAG compliance dashboard", "Community language & sign packs", "Certification programme"]),
]
cw = 3.95
for i, (t, accent, fill, items) in enumerate(phases):
    x = 0.7 + i * (cw + 0.27)
    rect(s, x, 1.65, cw, 4.5, fill, line=accent, line_w=2, rounded=True, radius=0.05, shadow=True)
    rect(s, x, 1.65, cw, 0.78, accent, rounded=True, radius=0.14); rect(s, x, 2.05, cw, 0.38, accent)
    txt(s, x+0.2, 1.65, cw-0.4, 0.78, t, size=16.5, color=CREAMT, bold=True, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    bullets(s, x+0.3, 2.7, cw-0.55, 3.3, items, size=14.5, gap=11, mcolor=accent)
    if i < 2: arrow(s, x+cw+0.02, 3.9, x+cw+0.25, 3.9, GOLD, 2.5)
txt(s, 0.7, 6.35, 12.0, 0.5, "*  Indian Sign Language avatar - our flagship next milestone.",
    size=13.5, color=MAROON, bold=True, italic=True)
footer(s)
notes(s, "Momentum + credible path. Working MVP today; the sign-language avatar is the headline "
          "next milestone; long term is shared state DPI with certification.")

# ============================================================ 18 ASK
s = slide(); title_bar(s, "The Ask")
rect(s, 0.7, 1.6, 12.0, 1.5, MAROON, rounded=True, radius=0.08, shadow=True)
txt(s, 0.9, 1.6, 11.6, 1.5, "Give us one department's website - we'll make it usable by every citizen.",
    size=24, color=CREAMT, bold=True, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
bullets(s, 1.0, 3.4, 11.4, 3.0, [
    ("A pilot:", "one Odisha government department or university portal to deploy on."),
    ("Access & mentorship:", "real portals to test against, and guidance to harden for production."),
    ("Support to scale:", "to build the sign-language avatar and a state-wide, open DPI rollout."),
], size=19, gap=18)
footer(s)
notes(s, "Make the ask concrete and small enough to say yes to: one pilot. Confidence line - hand "
          "us one site and we'll prove it.")

# ============================================================ 19 THANK YOU
s = slide()
rect(s, 0, 0, SW, SH, MAROON); rect(s, 0, 0, SW, 0.14, GOLD_B); rect(s, 0, SH-0.14, SW, 0.14, GOLD_B)
pic(s, AST / "chakra_gold.png", 5.97, 0.85, h=2.3)
txt(s, 0, 3.35, SW, 0.9, "Thank you", size=44, color=CREAMT, bold=True, align=PP_ALIGN.CENTER)
pic(s, AST / "odia_sub.png", (SW-6.0)/2, 4.35, h=0.5)
txt(s, 0, 5.0, SW, 0.5, "Accessibility as a Service  ·  Team SUBARNAREKHA", size=20, color=GOLD_B, bold=True, align=PP_ALIGN.CENTER)
txt(s, 0, 5.5, SW, 0.5, "Open-source · for the public good · accessibility for every citizen", size=15, color=CREAM, align=PP_ALIGN.CENTER)
notes(s, "Close warm and confident: open-source accessibility for every citizen, in their own "
          "language, on any site - and soon, in sign language. Invite the live demo.")

out = REPO / "pitch" / "AaaS-Pitch.pptx"
prs.save(str(out))
print("saved", out, "slides:", len(prs.slides._sldIdLst))
