"""Build the AaaS judge pitch deck (pitch/AaaS-Pitch.pptx).

16:9, Odisha/Indian theme (maroon-saffron-gold-cream, Konark chakra). Diagrams
are native editable shapes; Odia text + chakra are pre-rendered images so they
display on any machine. Each slide carries short speaker notes.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from copy import deepcopy
import pathlib

REPO = pathlib.Path(__file__).resolve().parents[1]
AST = REPO / "pitch" / "assets"

# palette
CREAM = RGBColor(0xFF, 0xF7, 0xE9)
PAPER = RGBColor(0xFD, 0xF3, 0xDD)
MAROON = RGBColor(0x8D, 0x1B, 0x2F)
MAROON_D = RGBColor(0x6E, 0x13, 0x22)
SAFFRON = RGBColor(0xEF, 0x7D, 0x16)
SAFF_L = RGBColor(0xFD, 0xEC, 0xD4)
GOLD = RGBColor(0xC9, 0x95, 0x2A)
GOLD_B = RGBColor(0xE6, 0xB9, 0x3F)
INDIGO = RGBColor(0x1B, 0x2A, 0x66)
IND_L = RGBColor(0xE9, 0xEC, 0xFB)
GREEN = RGBColor(0x1A, 0x7A, 0x3C)
GRN_L = RGBColor(0xE4, 0xF4, 0xE9)
INK = RGBColor(0x3A, 0x24, 0x10)
MUTED = RGBColor(0x7A, 0x5A, 0x38)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CREAMT = RGBColor(0xFF, 0xF3, 0xDA)

FONT = "Segoe UI"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = 13.333, 7.5
BLANK = prs.slide_layouts[6]


def slide():
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, SW, SH, CREAM)
    return s


def rect(s, x, y, w, h, fill, line=None, line_w=1.0, rounded=False, radius=0.09, shadow=False):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        try: shp.adjustments[0] = radius
        except Exception: pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    if shadow:
        el = shp._element.spPr
        ef = el.makeelement(qn('a:effectLst'), {}); el.append(ef)
        sh = ef.makeelement(qn('a:outerShdw'),
            {'blurRad':'90000','dist':'40000','dir':'5400000','rotWithShape':'0'}); ef.append(sh)
        c = sh.makeelement(qn('a:srgbClr'), {'val':'6E1322'}); sh.append(c)
        a = c.makeelement(qn('a:alpha'), {'val':'22000'}); c.append(a)
    return shp


def txt(s, x, y, w, h, runs, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, font=FONT, spacing=1.0, italic=False):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    items = runs if isinstance(runs, list) else [runs]
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing: p.line_spacing = spacing
        p.space_after = Pt(4)
        if isinstance(it, tuple):
            text, opts = it
        else:
            text, opts = it, {}
        r = p.add_run(); r.text = text
        f = r.font
        f.size = Pt(opts.get("size", size)); f.bold = opts.get("bold", bold)
        f.italic = opts.get("italic", italic); f.name = opts.get("font", font)
        f.color.rgb = opts.get("color", color)
    return tb


def bullets(s, x, y, w, h, items, size=18, color=INK, gap=6, marker="▸  ", mcolor=SAFFRON):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.05; p.space_after = Pt(gap)
        head = it[0] if isinstance(it, tuple) else None
        body = it[1] if isinstance(it, tuple) else it
        rm = p.add_run(); rm.text = marker
        rm.font.size = Pt(size); rm.font.bold = True; rm.font.name = FONT; rm.font.color.rgb = mcolor
        if head:
            rh = p.add_run(); rh.text = head + "  "
            rh.font.size = Pt(size); rh.font.bold = True; rh.font.name = FONT; rh.font.color.rgb = INK
        rb = p.add_run(); rb.text = body
        rb.font.size = Pt(size); rb.font.name = FONT; rb.font.color.rgb = color
    return tb


def arrow(s, x1, y1, x2, y2, color=GOLD, w=2.0):
    cn = s.shapes.add_connector(2, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    cn.line.color.rgb = color; cn.line.width = Pt(w)
    ln = cn.line._get_or_add_ln()
    end = ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'})
    ln.append(end)
    cn.shadow.inherit = False
    return cn


def pic(s, path, x, y, w=None, h=None):
    kw = {}
    if w: kw["width"] = Inches(w)
    if h: kw["height"] = Inches(h)
    return s.shapes.add_picture(str(path), Inches(x), Inches(y), **kw)


def title_bar(s, title, kicker=None):
    rect(s, 0, 0, SW, 1.0, MAROON)
    rect(s, 0, 1.0, SW, 0.07, GOLD_B)
    txt(s, 0.55, 0.0, 10.8, 1.0, title, size=30, color=CREAMT, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    if kicker:
        txt(s, 0.57, 0.62, 10.8, 0.35, kicker, size=14, color=GOLD_B, bold=True)
    pic(s, AST / "chakra_gold.png", 12.35, 0.13, h=0.74)


def footer(s, n):
    txt(s, 0.55, 7.06, 8.0, 0.35,
        "Team SUBARNAREKHA  ·  Smart Odisha Hackathon '25", size=10, color=MUTED)
    txt(s, 11.4, 7.06, 1.4, 0.35, f"{n:02d} / 16", size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def notes(s, text):
    s.notes_slide.notes_text_frame.text = text


def box(s, x, y, w, h, title, sub, fill, border, tcolor=INK, scolor=MUTED, tsize=15, ssize=11.5):
    rect(s, x, y, w, h, fill, line=border, line_w=1.75, rounded=True, radius=0.12)
    tb = s.shapes.add_textbox(Inches(x+0.12), Inches(y), Inches(w-0.24), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.05); tf.margin_right = Inches(0.05)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; p.line_spacing = 1.0
    r = p.add_run(); r.text = title; r.font.size = Pt(tsize); r.font.bold = True; r.font.name = FONT; r.font.color.rgb = tcolor
    if sub:
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER; p2.line_spacing = 1.0; p2.space_before = Pt(2)
        r2 = p2.add_run(); r2.text = sub; r2.font.size = Pt(ssize); r2.font.name = FONT; r2.font.color.rgb = scolor


# ============================================================ S1 TITLE
s = slide()
rect(s, 0, 0, SW, SH, MAROON)
rect(s, 0, 0, SW, 0.14, GOLD_B)
rect(s, 0, SH-0.14, SW, 0.14, GOLD_B)
pic(s, AST / "chakra_gold.png", 0.7, 1.5, h=3.4)
txt(s, 4.5, 1.5, 8.2, 1.2, "Accessibility as a Service", size=52, color=CREAMT, bold=True)
txt(s, 4.52, 2.55, 8.2, 0.6, "(AaaS)", size=26, color=GOLD_B, bold=True)
pic(s, AST / "odia_sub.png", 4.52, 3.2, h=0.62)
txt(s, 4.52, 3.95, 8.4, 0.6,
    "Shared accessibility infrastructure for public institutions", size=20, color=CREAM)
rect(s, 4.55, 4.7, 6.0, 0.04, GOLD_B)
txt(s, 4.5, 4.9, 8.4, 1.4, [
    ("Team SUBARNAREKHA", {"size": 24, "bold": True, "color": WHITE}),
    ("Smart Odisha Hackathon '25", {"size": 16, "color": GOLD_B}),
    ("Theme: Accessibility · E-commerce & Cyber Security   |   Software · Odiapreneur track",
     {"size": 13, "color": CREAM}),
], spacing=1.15)
notes(s, "Namaskar. We are Team SUBARNAREKHA. Our project is Accessibility as a Service - "
          "shared infrastructure that makes any government or education website usable by every "
          "citizen, in their own language, including the blind, low-vision and dyslexic. One line "
          "for institutions to integrate; one extension for citizens to use it everywhere.")

# ============================================================ S2 PROBLEM
s = slide(); title_bar(s, "The Problem")
bullets(s, 0.7, 1.45, 12.0, 5.4, [
    ("26.8M+ Indians live with a disability", "(Census 2011) - yet most government and education sites are unusable with a screen reader."),
    ("The web is English-first.", "Only about 1 in 10 Indians reads English comfortably - Odia-only citizens are shut out of digital services."),
    ("Today's \"fix\" is an overlay sold to site owners.", "It ignores Indian languages, only works where installed, and is widely criticised for not actually fixing access."),
    ("Every department builds accessibility alone.", "Slow, costly, inconsistent - and rarely WCAG / GIGW compliant."),
], size=19, gap=14)
rect(s, 0.7, 6.55, 12.0, 0.02, GOLD)
txt(s, 0.7, 6.62, 12.0, 0.5,
    "A citizen who cannot read English or cannot see the screen simply cannot use the service.",
    size=15, color=MAROON, bold=True, italic=True)
footer(s, 2)
notes(s, "Frame the gap: disability + language + no shared solution. The punchline - existing "
          "overlays are English-centric band-aids. Public digital services exclude exactly the "
          "people who need them most.")

# ============================================================ S3 WHO
s = slide(); title_bar(s, "Who gets left behind")
cards = [
    ("The blind citizen", "A government notice or scholarship form they literally cannot read - no audio, no screen-reader support.", MAROON, SAFF_L),
    ("The Odia speaker", "A farmer or student fluent in Odia, facing an English-only portal with no translation.", GREEN, GRN_L),
    ("The exam candidate", "A student with dyslexia in an online test - no audio support, no extra time, no friendly fonts.", INDIGO, IND_L),
]
cw, gap = 3.95, 0.30
x0 = 0.7
for i, (t, d, accent, fill) in enumerate(cards):
    x = x0 + i * (cw + gap)
    rect(s, x, 1.6, cw, 4.4, fill, line=accent, line_w=2, rounded=True, radius=0.06, shadow=True)
    rect(s, x, 1.6, cw, 0.7, accent, rounded=True, radius=0.18)
    rect(s, x, 1.95, cw, 0.35, accent)
    txt(s, x+0.25, 1.6, cw-0.5, 0.7, t, size=18, color=CREAMT, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, x+0.28, 2.6, cw-0.56, 3.2, d, size=15.5, color=INK, spacing=1.1)
txt(s, 0.7, 6.35, 12.0, 0.7,
    "These are not edge cases - they are millions of citizens that public services are meant to serve.",
    size=15, color=MAROON, bold=True, italic=True)
footer(s, 3)
notes(s, "Put a human face on the stats. Three personas the judges can picture. End on scale: "
          "millions, not edge cases.")

# ============================================================ S4 SOLUTION
s = slide(); title_bar(s, "Our Solution: one shared platform")
rect(s, 0.7, 1.45, 12.0, 1.15, MAROON, rounded=True, radius=0.12)
txt(s, 0.9, 1.45, 11.6, 1.15,
    "Accessibility as a Service - institutions plug in ONCE; citizens get it EVERYWHERE.",
    size=22, color=CREAMT, bold=True, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
pillars = [
    ("API + Widget", "Integrate once", "One <script> line or a REST API turns any site accessible.", SAFFRON, SAFF_L),
    ("Browser Extension", "Works on any site", "Citizens install it once - accessibility on every website, no site cooperation.", GREEN, GRN_L),
    ("On-device Indian AI", "Private & offline", "Neural Odia/Hindi/English voice & translation run in the browser.", INDIGO, IND_L),
]
cw = 3.95
for i, (t, tag, d, accent, fill) in enumerate(pillars):
    x = 0.7 + i * (cw + 0.30)
    rect(s, x, 2.95, cw, 3.5, fill, line=accent, line_w=2, rounded=True, radius=0.07, shadow=True)
    txt(s, x+0.28, 3.2, cw-0.56, 0.6, t, size=19, color=accent, bold=True)
    txt(s, x+0.28, 3.85, cw-0.56, 0.4, tag.upper(), size=12, color=MUTED, bold=True)
    txt(s, x+0.28, 4.35, cw-0.56, 2.0, d, size=15.5, color=INK, spacing=1.12)
footer(s, 4)
notes(s, "The one big idea: shared infrastructure, not another point tool. Two delivery modes "
          "(integrate vs extension) and the differentiator - Indian-language AI that runs on-device.")

# ============================================================ S5 TWO WAYS IN
s = slide(); title_bar(s, "Two ways in - the same powers")
left = [("Institutions integrate", INDIGO, IND_L, [
    "Add one widget <script> line, or call the REST API.",
    "Their site gains read-aloud, translation and voice-to-fill forms.",
    "Central control: API keys, usage metering, WCAG logging.",
])]
right = [("Citizens install the extension", GREEN, GRN_L, [
    "One-time install in Chrome / Edge.",
    "Accessibility on ANY website - including legacy portals that will never be updated.",
    "Runs on-device: offline and even on strict-CSP sites.",
])]
for col, data in ((0, left), (1, right)):
    t, accent, fill, items = data[0]
    x = 0.7 + col * 6.25
    rect(s, x, 1.55, 5.9, 4.4, fill, line=accent, line_w=2, rounded=True, radius=0.06, shadow=True)
    rect(s, x, 1.55, 5.9, 0.75, accent, rounded=True, radius=0.16)
    rect(s, x, 1.95, 5.9, 0.35, accent)
    txt(s, x+0.3, 1.55, 5.3, 0.75, t, size=19, color=CREAMT, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    bullets(s, x+0.32, 2.55, 5.3, 3.3, items, size=15.5, gap=10, mcolor=accent)
rect(s, 0.7, 6.35, 12.0, 0.7, GOLD_B, rounded=True, radius=0.3)
txt(s, 0.7, 6.35, 12.0, 0.7, "Power with the user - not just the site owner.",
    size=18, color=MAROON_D, bold=True, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
footer(s, 5)
notes(s, "This slide pre-empts the overlay question. Overlays are site-owner-only. We add a "
          "citizen-controlled path so accessibility reaches sites that will never integrate.")

# ============================================================ S6 ARCHITECTURE
s = slide(); title_bar(s, "Core architecture",
                       "two ways in -> one API gateway -> shared engines (also run on-device)")
for lx, lab in ((1.9, "CLIENTS"), (6.55, "API GATEWAY"), (10.95, "ENGINES")):
    txt(s, lx-1.4, 1.5, 2.8, 0.3, lab, size=12, color=MAROON, bold=True, align=PP_ALIGN.CENTER)
box(s, 0.7, 1.95, 2.95, 1.35, "Institution site / app", "widget (1 line) + REST API", SAFF_L, SAFFRON)
box(s, 0.7, 3.55, 2.95, 1.35, "Browser extension", "any website, on-device", SAFF_L, SAFFRON)
rect(s, 4.35, 1.95, 4.05, 2.95, MAROON, line=GOLD_B, line_w=2, rounded=True, radius=0.05, shadow=True)
txt(s, 4.45, 2.05, 3.85, 0.5, "Accessibility API Gateway", size=17, color=CREAMT, bold=True, align=PP_ALIGN.CENTER)
txt(s, 4.45, 2.5, 3.85, 0.3, "X-API-Key  ·  multi-tenant  ·  CORS", size=11.5, color=GOLD_B, align=PP_ALIGN.CENTER)
for i, ep in enumerate(["POST /tts", "POST /stt", "POST /translate", "GET /widget.js"]):
    cx = 4.5 + (i % 2) * 1.95; cy = 2.95 + (i // 2) * 0.5
    rect(s, cx, cy, 1.8, 0.4, CREAMT, rounded=True, radius=0.3)
    txt(s, cx, cy, 1.8, 0.4, ep, size=11, color=MAROON_D, bold=True, font="Consolas",
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
txt(s, 4.45, 4.05, 3.85, 0.3, "usage metering · rate limiting · WCAG logging",
    size=11, color=GOLD_B, align=PP_ALIGN.CENTER)
box(s, 4.35, 5.05, 4.05, 0.95, "Admin Dashboard", "tenants · API keys · usage · WCAG", INDIGO, GOLD_B, tcolor=CREAMT, scolor=RGBColor(0xBC,0xD0,0xFF))
arrow(s, 4.35, 5.05, 6.37, 4.95, GOLD)
box(s, 9.05, 1.95, 3.55, 0.85, "Text-to-Speech + Screen Reader", "Meta MMS-TTS · Odia/Hindi/English", IND_L, INDIGO, tsize=13.5, ssize=11)
box(s, 9.05, 2.98, 3.55, 0.85, "Speech-to-Text", "Whisper", SAFF_L, SAFFRON, tsize=13.5, ssize=11)
box(s, 9.05, 4.01, 3.55, 0.85, "Translation", "AI4Bharat IndicTrans2", GRN_L, GREEN, tsize=13.5, ssize=11)
arrow(s, 3.65, 2.62, 4.35, 2.9, GOLD)
arrow(s, 3.65, 4.22, 4.35, 3.6, GOLD)
arrow(s, 8.4, 2.7, 9.05, 2.38, GOLD)
arrow(s, 8.4, 3.2, 9.05, 3.4, GOLD)
arrow(s, 8.4, 3.7, 9.05, 4.43, GOLD)
txt(s, 0.7, 6.25, 12.0, 0.7,
    "On-device option: the extension runs all three engines in the browser via WebAssembly - "
    "offline, and even on strict-CSP sites.", size=13.5, color=MUTED, italic=True)
footer(s, 6)
notes(s, "Walk left to right. The gateway is the product: secure, multi-tenant, metered. The same "
          "engines run server-side OR on-device in the extension. Admin dashboard gives departments "
          "oversight and WCAG evidence.")

# ============================================================ S7 HOW IT WORKS
s = slide(); title_bar(s, "How it works")
rows = [
    ("Read aloud", INDIGO, IND_L, ["Page text", "IndicTrans2 (if needed)", "MMS-TTS on-device", "Odia voice plays"]),
    ("Translate page", GREEN, GRN_L, ["Page text", "IndicTrans2", "Page rewritten in Odia"]),
    ("Speak to fill", SAFFRON, SAFF_L, ["Voice (mic)", "Whisper on-device", "Form field auto-filled"]),
]
y = 1.7
for name, accent, fill, steps in rows:
    rect(s, 0.7, y, 2.3, 0.95, accent, rounded=True, radius=0.12)
    txt(s, 0.75, y, 2.2, 0.95, name, size=16, color=CREAMT, bold=True, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    sx = 3.25; sw = (12.6 - sx) / len(steps) - 0.25
    for i, st in enumerate(steps):
        x = sx + i * (sw + 0.25)
        rect(s, x, y+0.12, sw, 0.7, fill, line=accent, line_w=1.5, rounded=True, radius=0.12)
        txt(s, x+0.05, y+0.12, sw-0.1, 0.7, st, size=12.5, color=INK, bold=True,
            anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
        if i < len(steps)-1:
            arrow(s, x+sw, y+0.47, x+sw+0.25, y+0.47, accent, 1.75)
    y += 1.35
rect(s, 0.7, 6.25, 12.0, 0.8, PAPER, line=GOLD, line_w=1.25, rounded=True, radius=0.2)
txt(s, 0.95, 6.25, 11.5, 0.8,
    "\"On-device\" = runs inside the browser via WebAssembly. No server needed, fully offline, "
    "and translation is skipped when the page is already in your language.",
    size=13.5, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
footer(s, 7)
notes(s, "Three everyday actions, each a small pipeline. Stress that the heavy AI runs on the "
          "user's device - so it works with no backend and respects privacy.")

# ============================================================ S8 DEMO
s = slide(); title_bar(s, "It already works - live, on any site")
pic(s, AST / "demo-wikipedia.png", 0.7, 1.45, w=8.7)
rect(s, 0.7, 1.45, 8.7, 4.13, None, line=GOLD, line_w=1.5)
txt(s, 9.65, 1.7, 3.2, 4.2, [
    ("English Wikipedia -> Odia", {"size": 18, "bold": True, "color": MAROON}),
    ("Title and page flip to Odia in place, and the page reads aloud in a neural Odia voice.",
     {"size": 14, "color": INK}),
    ("No backend running.", {"size": 14, "bold": True, "color": GREEN}),
    ("Works despite Wikipedia's strict Content-Security-Policy - because translation goes through "
     "the extension and the voice runs on-device.", {"size": 14, "color": INK}),
], spacing=1.12)
footer(s, 8)
notes(s, "This is a real screenshot, not a mockup. Wikipedia is a deliberately hard target "
          "(strict CSP). If it works here, it works on government portals. Offer to do it live.")

# ============================================================ S9 FIVE PARTS
s = slide(); title_bar(s, "One platform, five parts")
parts = [
    ("1", "Speech-to-Text", "Voice input, dictation", SAFFRON),
    ("2", "Text-to-Speech + Screen Reader", "Neural Odia/Hindi/English", INDIGO),
    ("3", "Accessibility API Gateway", "The plug-and-play surface", MAROON),
    ("4", "Translation - IndicTrans2", "Indic-tuned, Odia-first", GREEN),
    ("5", "Admin Dashboard", "Tenants, keys, usage, WCAG", INDIGO),
]
cw, ch = 3.95, 1.55
for i, (n, t, d, accent) in enumerate(parts):
    x = 0.7 + (i % 3) * (cw + 0.27)
    y = 1.7 + (i // 3) * (ch + 0.3)
    rect(s, x, y, cw, ch, PAPER, line=GOLD, line_w=1.25, rounded=True, radius=0.08, shadow=True)
    rect(s, x+0.2, y+0.3, 0.95, 0.95, accent, rounded=True, radius=0.5)
    txt(s, x+0.2, y+0.3, 0.95, 0.95, n, size=30, color=CREAMT, bold=True, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    txt(s, x+1.35, y+0.32, cw-1.55, 0.95, [
        (t, {"size": 15.5, "bold": True, "color": INK}),
        (d, {"size": 12.5, "color": MUTED}),
    ], spacing=1.05)
footer(s, 9)
notes(s, "Maps directly to the problem statement's modules. We didn't pick one piece - we "
          "built the whole accessibility platform, with admin oversight and WCAG logging as "
          "differentiators.")

# ============================================================ S10 TECH
s = slide(); title_bar(s, "Built to run anywhere")
bullets(s, 0.7, 1.5, 7.4, 5.2, [
    ("Browser extension (MV3):", "MAIN + isolated worlds, a Shadow-DOM widget that can't clash with any page."),
    ("FastAPI gateway:", "multi-tenant, X-API-Key auth, CORS, usage metering, WCAG logging."),
    ("Voices - Meta MMS-TTS:", "Odia, Hindi & English neural speech."),
    ("Speech-to-Text - Whisper;", "Translation - AI4Bharat IndicTrans2."),
    ("On-device runtime:", "transformers.js + ONNX / WebAssembly - inference in the browser."),
    ("Portable bundle:", "the whole stack runs offline from a USB / folder (start.bat)."),
], size=16.5, gap=11)
rect(s, 8.4, 1.6, 4.25, 4.9, MAROON, rounded=True, radius=0.06, shadow=True)
txt(s, 8.7, 1.85, 3.7, 0.5, "Why it matters", size=18, color=GOLD_B, bold=True)
bullets(s, 8.7, 2.55, 3.7, 3.7, [
    "Privacy: voice & text stay on the device.",
    "Offline: no connectivity required.",
    "Resilient: works on strict-CSP sites.",
    "Low cost: no per-call cloud AI bills.",
], size=14.5, gap=12, mcolor=GOLD_B, color=CREAM)
footer(s, 10)
notes(s, "For technical judges. The headline engineering win is on-device inference: privacy, "
          "offline, resilience and cost - all at once. Mention the Shadow DOM isolation so the "
          "widget never breaks host pages.")

# ============================================================ S11 OVERLAY
s = slide(); title_bar(s, "Isn't this just another accessibility overlay?")
rows = [
    ("", "Overlays (accessiBe / UserWay)", "AaaS"),
    ("Who controls it", "Site owner installs & pays; one site only", "User installs the extension -> works on ANY site"),
    ("Language", "English / Western-centric", "Odia & Hindi first, via IndicTrans2"),
    ("Data", "Cloud SaaS; content sent to their servers", "On-device & in-country; DPDP-friendly"),
    ("Claim", "\"Auto-makes you compliant\" (criticised)", "Honest: assistive layer + API toolkit"),
    ("Extras", "Generic toolbar", "Voice-to-fill, admin & WCAG logs"),
]
tw, th = 12.0, 4.9
tbl = s.shapes.add_table(len(rows), 3, Inches(0.7), Inches(1.5), Inches(tw), Inches(th)).table
tbl.columns[0].width = Inches(2.4); tbl.columns[1].width = Inches(4.8); tbl.columns[2].width = Inches(4.8)
for r, row in enumerate(rows):
    tbl.rows[r].height = Inches(th/len(rows))
    for c, val in enumerate(row):
        cell = tbl.cell(r, c); cell.margin_left = Inches(0.12); cell.margin_right = Inches(0.1)
        cell.margin_top = Inches(0.04); cell.margin_bottom = Inches(0.04)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame; tf.word_wrap = True; p = tf.paragraphs[0]
        rn = p.add_run(); rn.text = val; rn.font.name = FONT
        if r == 0:
            cell.fill.solid(); cell.fill.fore_color.rgb = MAROON if c != 2 else GREEN
            rn.font.color.rgb = CREAMT; rn.font.bold = True; rn.font.size = Pt(14.5)
        else:
            cell.fill.solid(); cell.fill.fore_color.rgb = PAPER if c != 2 else GRN_L
            rn.font.size = Pt(13)
            rn.font.color.rgb = INK
            if c == 0: rn.font.bold = True; rn.font.color.rgb = MAROON
            if c == 2: rn.font.bold = True; rn.font.color.rgb = RGBColor(0x0F,0x5A,0x2C)
txt(s, 0.7, 6.55, 12.0, 0.5,
    "We are not a compliance shortcut - we are accessibility infrastructure with first-class Indian-language access.",
    size=14, color=MAROON, bold=True, italic=True)
footer(s, 11)
notes(s, "The expected hard question - answer it head-on. Key lines: overlays are site-owner, "
          "English, cloud, and over-claim compliance. We are user-controlled, Indian-language, "
          "on-device, and honest about being an assistive layer + toolkit.")

# ============================================================ S12 WHY GOVT
s = slide(); title_bar(s, "Why government should adopt AaaS")
bullets(s, 0.7, 1.5, 12.0, 5.4, [
    ("Compliance, built in:", "helps meet the RPwD Act 2016, GIGW 3.0 and WCAG 2.1 - with WCAG result logging for evidence."),
    ("Real inclusion:", "first-class Odia & Hindi reaches the citizens English-only tools miss."),
    ("Data sovereignty:", "on-device & in-country processing - DPDP Act 2023 friendly; no citizen data sent to foreign clouds."),
    ("Cost-efficient:", "one shared platform instead of every department reinventing it - and it works on existing legacy portals, no rebuild."),
    ("Maximum reach:", "the extension covers sites that cannot be changed today, so impact is immediate."),
], size=18, gap=15)
footer(s, 12)
notes(s, "Speak the government's language: compliance (RPwD/GIGW/WCAG), sovereignty (DPDP, "
          "on-device), cost (shared infra, no rebuild) and reach (works on legacy sites now). "
          "These are procurement decision factors.")

# ============================================================ S13 IMPACT
s = slide(); title_bar(s, "Impact & alignment")
rect(s, 0.7, 1.55, 5.9, 4.9, PAPER, line=GOLD, line_w=1.25, rounded=True, radius=0.05, shadow=True)
txt(s, 0.95, 1.75, 5.4, 0.5, "Outcomes", size=18, color=MAROON, bold=True)
bullets(s, 0.95, 2.4, 5.4, 3.8, [
    "Any citizen can hear & read any public page in Odia.",
    "Forms get filled by voice - no typing, no scribe needed.",
    "Departments get measurable WCAG compliance.",
    "Accessibility delivered without rebuilding old sites.",
], size=15.5, gap=12)
rect(s, 6.75, 1.55, 5.9, 4.9, MAROON, rounded=True, radius=0.05, shadow=True)
txt(s, 7.0, 1.75, 5.4, 0.5, "Aligns with national missions", size=18, color=GOLD_B, bold=True)
bullets(s, 7.0, 2.4, 5.4, 3.8, [
    "Digital India - inclusive e-governance.",
    "Sugamya Bharat (Accessible India Campaign).",
    "SDG 10 - reduced inequalities.",
    "SDG 4 - inclusive, quality education.",
], size=15.5, gap=12, mcolor=GOLD_B, color=CREAM)
footer(s, 13)
notes(s, "Tie outcomes to flagship missions judges know. This signals the project isn't a toy - "
          "it advances policy goals the state already cares about.")

# ============================================================ S14 ROADMAP
s = slide(); title_bar(s, "Way forward")
phases = [
    ("NOW - MVP", GREEN, GRN_L, ["Extension + gateway live", "Odia / Hindi / English", "On-device TTS & STT", "Offline portable bundle"]),
    ("NEXT - 3-6 months", SAFFRON, SAFF_L, ["More Indian languages", "Pilot with one govt dept", "Android / mobile", "Sign-language avatar (R&D)"]),
    ("LATER - 6-12 months", INDIGO, IND_L, ["State-wide rollout", "WCAG compliance dashboard", "Developer SDKs & marketplace", "Certification programme"]),
]
cw = 3.95
for i, (t, accent, fill, items) in enumerate(phases):
    x = 0.7 + i * (cw + 0.27)
    rect(s, x, 1.7, cw, 4.6, fill, line=accent, line_w=2, rounded=True, radius=0.05, shadow=True)
    rect(s, x, 1.7, cw, 0.8, accent, rounded=True, radius=0.14)
    rect(s, x, 2.1, cw, 0.4, accent)
    txt(s, x+0.2, 1.7, cw-0.4, 0.8, t, size=16.5, color=CREAMT, bold=True, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    bullets(s, x+0.3, 2.75, cw-0.55, 3.3, items, size=15, gap=12, mcolor=accent)
    if i < 2:
        arrow(s, x+cw+0.02, 4.0, x+cw+0.25, 4.0, GOLD, 2.5)
footer(s, 14)
notes(s, "Show momentum and a credible path. We already have a working MVP; the near-term ask is "
          "a real pilot; the long term is shared state infrastructure with certification.")

# ============================================================ S15 ASK
s = slide(); title_bar(s, "The Ask")
rect(s, 0.7, 1.6, 12.0, 1.5, MAROON, rounded=True, radius=0.08, shadow=True)
txt(s, 0.9, 1.6, 11.6, 1.5,
    "Give us one department's website - we'll make it usable by every citizen.",
    size=24, color=CREAMT, bold=True, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
bullets(s, 1.0, 3.4, 11.4, 3.0, [
    ("A pilot:", "one Odisha government department or university portal to deploy on."),
    ("Access & mentorship:", "real portals to test against, and guidance to harden for production."),
    ("Support to scale:", "security hardening and infrastructure for a state-wide rollout."),
], size=19, gap=18)
footer(s, 15)
notes(s, "Make the ask concrete and small enough to say yes to: one pilot. Confidence line - hand "
          "us one site and we will prove it.")

# ============================================================ S16 THANK YOU
s = slide()
rect(s, 0, 0, SW, SH, MAROON)
rect(s, 0, 0, SW, 0.14, GOLD_B); rect(s, 0, SH-0.14, SW, 0.14, GOLD_B)
pic(s, AST / "chakra_gold.png", 5.97, 0.95, h=2.4)
txt(s, 0, 3.5, SW, 0.9, "Thank you", size=46, color=CREAMT, bold=True, align=PP_ALIGN.CENTER)
pic(s, AST / "odia_sub.png", (SW-6.0)/2, 4.5, h=0.5)
txt(s, 0, 5.15, SW, 0.5, "Accessibility as a Service  ·  Team SUBARNAREKHA",
    size=20, color=GOLD_B, bold=True, align=PP_ALIGN.CENTER)
txt(s, 0, 5.7, SW, 0.5, "Smart Odisha Hackathon '25  ·  Accessibility for every citizen",
    size=15, color=CREAM, align=PP_ALIGN.CENTER)
notes(s, "Close warm and confident. Restate the one-liner: accessibility for every citizen, in "
          "their own language, on any site. Invite the live demo.")

out = REPO / "pitch" / "AaaS-Pitch.pptx"
prs.save(str(out))
print("saved", out, "slides:", len(prs.slides._sldIdLst))
